import csv
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from jsonschema import Draft7Validator, FormatChecker


MAX_ASSET_FILE_BYTES = 5 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}
CIAAAN_COLUMNS = (
    "confidentiality",
    "integrity",
    "availability",
    "authenticity",
    "authorization",
    "non_repudiation",
)
REQUIRED_COLUMNS = {
    "asset_title",
    "asset_type",
    "asset_description",
    *CIAAAN_COLUMNS,
}
ALLOWED_COLUMNS = REQUIRED_COLUMNS | {"asset_id"}
ASSET_TYPES = {
    "communication_path",
    "data_store",
    "ecu",
    "function",
    "auth_credential",
    "api_endpoint",
    "cloud_service",
}


class AssetImportError(ValueError):
    pass


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def _parse_boolean(value: Any, row_number: int, column: str) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in {0, 1}:
        return bool(value)
    normalized = str(value or "").strip().lower()
    if normalized in {"true", "yes", "y", "1"}:
        return True
    if normalized in {"false", "no", "n", "0"}:
        return False
    raise AssetImportError(
        f"Row {row_number}, {column}: expected true/false, yes/no, or 1/0"
    )


def _csv_rows(content: bytes) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AssetImportError("CSV must use UTF-8 encoding") from exc
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        raise AssetImportError("The file has no header row")
    headers = [_normalize_header(value) for value in reader.fieldnames]
    if len(headers) != len(set(headers)):
        raise AssetImportError("The header row contains duplicate columns")
    rows = []
    for row_number, raw in enumerate(reader, start=2):
        if None in raw:
            raise AssetImportError(f"Row {row_number}: contains more values than the header row")
        row = {_normalize_header(key): value for key, value in raw.items()}
        if any(str(value or "").strip() for value in row.values()):
            rows.append(row)
    return rows


def _xlsx_rows(content: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl exposes several format-specific exceptions
        raise AssetImportError("The XLSX workbook could not be read") from exc
    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        header_row = next(values, None)
        if header_row is None:
            raise AssetImportError("The workbook has no header row")
        headers = [_normalize_header(value) for value in header_row]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or any(not header for header in headers):
            raise AssetImportError("The header row contains a blank column")
        if len(headers) != len(set(headers)):
            raise AssetImportError("The header row contains duplicate columns")
        rows = []
        for values_row in values:
            row = {
                header: values_row[index] if index < len(values_row) else None
                for index, header in enumerate(headers)
            }
            if any(str(value or "").strip() for value in row.values()):
                rows.append(row)
        return rows
    finally:
        workbook.close()


def _schema() -> dict[str, Any]:
    repo_root = Path(__file__).resolve().parents[2]
    path = repo_root / "src" / "schemas" / "stage-03-asset-register.schema.json"
    return json.loads(path.read_text(encoding="utf-8"))


def import_asset_register(filename: str, content: bytes) -> list[dict[str, Any]]:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise AssetImportError("File must be .csv or .xlsx")
    if not content:
        raise AssetImportError("The uploaded file is empty")
    if len(content) > MAX_ASSET_FILE_BYTES:
        raise AssetImportError("File too large - max 5MB")

    rows = _csv_rows(content) if extension == ".csv" else _xlsx_rows(content)
    if not rows:
        raise AssetImportError("The file contains no Asset rows")

    columns = set(rows[0])
    missing = sorted(REQUIRED_COLUMNS - columns)
    unexpected = sorted(columns - ALLOWED_COLUMNS)
    if missing:
        raise AssetImportError(f"Missing required columns: {', '.join(missing)}")
    if unexpected:
        raise AssetImportError(f"Unexpected columns: {', '.join(unexpected)}")

    timestamp = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    assets = []
    seen_ids: set[str] = set()
    for index, row in enumerate(rows, start=2):
        asset_id = str(row.get("asset_id") or f"AS_{index - 1:02d}").strip()
        if not re.fullmatch(r"AS_\d{2,}", asset_id):
            raise AssetImportError(f"Row {index}, asset_id: expected AS_ followed by at least two digits")
        if asset_id in seen_ids:
            raise AssetImportError(f"Row {index}, asset_id: duplicate value {asset_id}")
        seen_ids.add(asset_id)

        asset_type = str(row.get("asset_type") or "").strip()
        if asset_type not in ASSET_TYPES:
            raise AssetImportError(
                f"Row {index}, asset_type: {asset_type!r} is not an allowed Asset type"
            )
        ciaaan = {
            column: _parse_boolean(row.get(column), index, column)
            for column in CIAAAN_COLUMNS
        }
        assets.append({
            "asset_id": asset_id,
            "asset_title": str(row.get("asset_title") or "").strip(),
            "asset_type": asset_type,
            "asset_description": str(row.get("asset_description") or "").strip(),
            "ciaaan": ciaaan,
            "input_mode": "manual",
            "created_timestamp": timestamp,
        })

    validator = Draft7Validator(_schema(), format_checker=FormatChecker())
    errors = sorted(validator.iter_errors(assets), key=lambda error: list(error.path))
    if errors:
        error = errors[0]
        path = list(error.path)
        if path and isinstance(path[0], int):
            row_number = path.pop(0) + 2
            column = ".".join(str(part) for part in path)
            prefix = f"Row {row_number}, {column}: " if column else f"Row {row_number}: "
        else:
            location = ".".join(str(part) for part in path)
            prefix = f"Invalid Asset data at {location}: " if location else "Invalid Asset data: "
        raise AssetImportError(f"{prefix}{error.message}")
    return assets
